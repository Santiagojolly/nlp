"""Extract text from plain-text, PDF, and DOCX documents."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}


class DocumentParseError(ValueError):
    """Raised when a document is unsupported or cannot be read."""


def extract_text(file_name: str, content: bytes) -> str:
    """Return extracted document text, raising a helpful error on failure."""
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentParseError(f"Unsupported file type. Choose: {supported}.")
    if not content:
        return ""

    try:
        if extension == ".txt":
            try:
                return content.decode("utf-8-sig")
            except UnicodeDecodeError:
                return content.decode("cp1252")

        if extension == ".pdf":
            reader = PdfReader(BytesIO(content), strict=False)
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)

        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                "\t".join(cell.text for cell in row.cells) for row in table.rows
            )
        return "\n".join(paragraphs)
    except Exception as exc:
        raise DocumentParseError(
            f"Could not extract text from {file_name}: {exc}"
        ) from exc