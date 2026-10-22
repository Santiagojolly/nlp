from io import BytesIO

import pytest
from docx import Document

from src.document_parser import DocumentParseError, extract_text


def test_extracts_utf8_text_files():
    assert extract_text("notes.txt", "Caf\u00e9 notes".encode("utf-8")) == "Caf\u00e9 notes"


def test_extracts_docx_paragraphs_and_table_cells():
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Document paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "First cell"
    table.cell(0, 1).text = "Second cell"
    document.save(buffer)

    extracted = extract_text("notes.docx", buffer.getvalue())

    assert "Document paragraph" in extracted
    assert "First cell\tSecond cell" in extracted


def test_rejects_unsupported_extensions_and_malformed_pdfs():
    with pytest.raises(DocumentParseError, match="Unsupported file type"):
        extract_text("notes.rtf", b"text")
    with pytest.raises(DocumentParseError, match="Could not extract text"):
        extract_text("broken.pdf", b"not a PDF")